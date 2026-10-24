from __future__ import annotations

import csv
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

import fitz
from openpyxl import Workbook

from .engine import page_range
from .storage import Workspace, app_root, safe_name
from .processes import run_converter


def office_binary():
    root = app_root() / "vendor" / "libreoffice"
    configured = os.environ.get("PDFSTUDIO_SOFFICE")
    candidates = [Path(configured)] if configured else []
    candidates += [root / "program" / "soffice.com", root / "program" / "soffice.exe"]
    if root.exists():
        candidates += list(root.glob("*/program/soffice.com")) + list(root.glob("**/program/soffice.exe"))
    available = shutil.which("soffice")
    if available:
        candidates.append(Path(available))
    return next((p for p in candidates if p.is_file()), None)


def office_to_pdf(source: Path, destination: Path, *, cancel=None):
    binary = office_binary()
    if not binary:
        raise ValueError("The local Office converter is not bundled yet. Run the runtime preparation step; no Microsoft Office is required.")
    destination.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="pdfstudio-office-") as profile:
        # Use an isolated converter profile with executable/linked content disabled.
        user = Path(profile) / "user"
        user.mkdir()
        (user / "registrymodifications.xcu").write_text('''<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting">
<prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop>
<prop oor:name="DisableActiveContent" oor:op="fuse"><value>true</value></prop>
<prop oor:name="BlockUntrustedRefererLinks" oor:op="fuse"><value>true</value></prop>
</item></oor:items>''', encoding="utf-8")
        result = run_converter([str(binary), f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(destination), str(source)], timeout=240, cancel=cancel)
    output = destination / (source.stem + ".pdf")
    if result.returncode or not output.is_file():
        raise ValueError("The document could not be converted by the local Office converter. Check its format or password protection.")
    return output


def extract_tables(workspace: Workspace, key: str, pages="all", strategy="lines"):
    if strategy not in {"lines", "lines_strict", "text"}:
        raise ValueError("Choose ruled or text-aligned table detection.")
    tables = []
    with fitz.open(workspace.document(key).path) as pdf:
        for index in page_range(pages, len(pdf)):
            page = pdf[index]
            found = page.find_tables(strategy=strategy)
            for i, table in enumerate(found.tables):
                tables.append({"page": index + 1, "table": i + 1, "cells": [[str(c) if c is not None else "" for c in row] for row in table.extract()]})
    return tables


def tables_export(tables: list[dict], format="xlsx"):
    if not tables:
        raise ValueError("No tables were detected. Try text-aligned detection, another page range, or use a PDF with selectable text.")
    out = io.BytesIO()
    if format == "xlsx":
        wb = Workbook()
        summary = wb.active
        summary.title = "Summary"
        summary.append(["Page", "Table", "Rows", "Columns"])
        for item in tables:
            rows = item["cells"]
            summary.append([item["page"], item["table"], len(rows), max((len(r) for r in rows), default=0)])
            ws = wb.create_sheet(f"P{item['page']}_T{item['table']}")
            for row in rows:
                # Treat extracted strings as data, never executable spreadsheet formulas.
                ws.append(["'" + c if str(c).startswith(("=", "+", "-", "@")) else c for c in row])
            ws.freeze_panes = "A2"
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in column) + 2))
        wb.save(out)
    else:
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
            for item in tables:
                text = io.StringIO(newline="")
                csv.writer(text).writerows(item["cells"])
                archive.writestr(f"page_{item['page']}_table_{item['table']}.csv", text.getvalue().encode("utf-8-sig"))
            archive.writestr("summary.json", json.dumps([{k: v for k, v in t.items() if k != "cells"} for t in tables], indent=2))
    return out.getvalue()


def markdown(workspace: Workspace, key: str, pages="all"):
    parts = []
    with fitz.open(workspace.document(key).path) as pdf:
        for index in page_range(pages, len(pdf)):
            page = pdf[index]
            blocks = page.get_text("dict", sort=True)["blocks"]
            sizes = Counter(round(span["size"]) for b in blocks if b["type"] == 0 for line in b["lines"] for span in line["spans"] if span["text"].strip())
            body = sizes.most_common(1)[0][0] if sizes else 12
            parts.append(f"<!-- Page {index + 1} -->")
            for block in blocks:
                if block["type"] != 0:
                    continue
                text = "\n".join("".join(span["text"] for span in line["spans"]) for line in block["lines"]).strip()
                if not text:
                    continue
                largest = max((span["size"] for line in block["lines"] for span in line["spans"]), default=body)
                prefix = "# " if largest >= body * 1.6 else "## " if largest >= body * 1.25 else ""
                parts.append(prefix + text + "\n")
    value = "\n\n".join(parts)
    if not value.replace("<!--", "").strip() or not any(p for p in parts if not p.startswith("<!--")):
        raise ValueError("No extractable text was found. Image-only scans need a PDF with selectable text.")
    return value


def payment_fields(path: Path, name: str):
    with fitz.open(path) as pdf:
        text = "\n".join(p.get_text() for p in pdf)
        words = [word for page in pdf for word in page.get_text("words")]
    if "Transaction summary" not in text:
        raise ValueError(f"{name} does not look like an Upwork transaction summary.")
    fields = {"OriginalName": re.sub(r"^\d{4}-\d{2}-\d{2}_[a-f0-9]+-", "", Path(name).stem)}
    patterns = {"ReferenceID": r"Reference ID\s*#\s*(\d+)", "PaymentAmount": r"Attached\s+[^\d]*([\d,.]+)", "Type": r"\b(Fixed-price|Hourly|Consultation|Bonus)\b", "ContractAmount": r"Contract:\s*([^\s]+)", "Fees": r"Fees:\s*([^\s]+)", "Taxes": r"Taxes:\s*([^\s]+)"}
    for key, pattern in patterns.items():
        match = re.search(pattern, text)
        if match:
            fields[key] = match[1]
    found = re.search(r"([A-Za-z]+ \d{1,2}, \d{4})", text)
    if found:
        for fmt in ["%b %d, %Y", "%B %d, %Y"]:
            try:
                fields["Date"] = datetime.strptime(found[1], fmt).strftime("%Y%m%d")
                break
            except ValueError:
                pass
    headers = [w for w in words if w[4] == "Freelancer"]
    if headers:
        header = headers[0]
        candidates = [w for w in words if w[1] > header[1] + 4 and w[0] >= header[0] - 10]
        if candidates:
            top = min(w[1] for w in candidates)
            row = sorted([w for w in candidates if abs(w[1] - top) < 5], key=lambda w: w[0])
            fields["Freelancer"] = " ".join(w[4] for w in row if not re.fullmatch(r"[+-]?\$[\d,.]+", w[4]))
            fields["ChargeAmount"] = " ".join(w[4] for w in row if re.fullmatch(r"[+-]?\$[\d,.]+", w[4]))
    return fields


def rename_preview(workspace: Workspace, ids: list[str], pattern: str):
    if not pattern.strip():
        raise ValueError("Enter a filename pattern.")
    result, used = [], set()
    for key in ids:
        doc = workspace.document(key)
        fields = payment_fields(doc.path, doc.name)
        required = re.findall(r"\{([^{}]+)\}", pattern)
        missing = [field for field in required if not str(fields.get(field, "")).strip()]
        if missing:
            raise ValueError(f"{doc.name}: missing fields {', '.join(missing)}. Change the pattern or inspect the document.")
        name = pattern
        for field in required:
            name = name.replace("{" + field + "}", str(fields[field]))
        name = safe_name(name.removesuffix(".pdf"))
        candidate, suffix = name + ".pdf", 2
        while candidate.casefold() in used:
            candidate = f"{name}_{suffix}.pdf"
            suffix += 1
        used.add(candidate.casefold())
        result.append({"id": key, "original": doc.name, "name": candidate, "fields": fields})
    return result
