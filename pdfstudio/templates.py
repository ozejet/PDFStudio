from __future__ import annotations

import copy
import csv
import io
import re
import uuid
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

import fitz
from openpyxl import load_workbook

from .engine import edit, timestamp
from .storage import Workspace, safe_name


def spreadsheet(data: bytes, name: str, sheet: str | None = None) -> dict:
    if name.lower().endswith(".csv"):
        text = data.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text))
        raw = list(reader)
        sheets = ["CSV"]
    else:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            sheets = workbook.sheetnames
            ws = workbook[sheet] if sheet else workbook.active
            if ws.max_row > 50001 or ws.max_column > 500:
                raise ValueError("Use a spreadsheet with at most 50,000 data rows and 500 columns.")
            raw = []
            for row in ws.iter_rows():
                values = []
                for cell in row:
                    value = cell.value
                    if isinstance(value, (date, datetime)):
                        value = value.isoformat()
                    elif isinstance(value, (int, float)) and re.fullmatch(r"0{2,}", cell.number_format or ""):
                        value = str(int(value)).zfill(len(cell.number_format))
                    values.append(value if value is not None else "")
                raw.append(values)
        finally:
            workbook.close()
    if not raw:
        raise ValueError("The spreadsheet is empty.")
    headers = [str(v).strip() if v is not None else "" for v in raw[0]]
    indices = [i for i, h in enumerate(headers) if h]
    columns = [headers[i] for i in indices]
    if not columns or len(columns) != len(set(columns)):
        raise ValueError("The header row must contain unique, non-empty column names.")
    rows = [{headers[i]: row[i] if i < len(row) else "" for i in indices} for row in raw[1:] if any(str(v).strip() for v in row if v is not None)]
    if len(rows) > 50000:
        raise ValueError("Use a spreadsheet with at most 50,000 data rows.")
    return {"columns": columns, "rows": rows, "sheets": sheets, "name": safe_name(name)}


THAI_DIGITS = ["ศูนย์", "หนึ่ง", "สอง", "สาม", "สี่", "ห้า", "หก", "เจ็ด", "แปด", "เก้า"]


def thai_integer(number: int):
    if number == 0:
        return THAI_DIGITS[0]
    if number >= 1_000_000:
        quotient, remainder = divmod(number, 1_000_000)
        return thai_integer(quotient) + "ล้าน" + (thai_integer(remainder) if remainder else "")
    result = []
    s = str(number)
    for i, char in enumerate(s):
        digit = int(char)
        position = len(s) - i - 1
        if not digit:
            continue
        if position == 1:
            result.append("สิบ" if digit == 1 else "ยี่สิบ" if digit == 2 else THAI_DIGITS[digit] + "สิบ")
        elif position == 0 and digit == 1 and number > 1:
            result.append("เอ็ด")
        else:
            result.append(THAI_DIGITS[digit] + ["", "สิบ", "ร้อย", "พัน", "หมื่น", "แสน"][position])
    return "".join(result)


def baht_text(value):
    try:
        amount = Decimal(str(value).replace(",", "")).quantize(Decimal(".01"), rounding=ROUND_HALF_UP)
        if not amount.is_finite():
            raise InvalidOperation
    except InvalidOperation:
        raise ValueError(f"Cannot turn {value!r} into Thai baht text.") from None
    negative = amount < 0
    amount = abs(amount)
    whole = int(amount)
    fraction = int((amount - whole) * 100)
    return ("ลบ" if negative else "") + thai_integer(whole) + "บาท" + (thai_integer(fraction) + "สตางค์" if fraction else "ถ้วน")


def resolve(item: dict, row: dict, index=1):
    source = item.get("sourceType", "constant")
    column = item.get("column", "")
    if source in {"excel", "computed"}:
        if column and column not in row:
            raise ValueError(f"Missing Excel column: {column}.")
        value = row.get(column, "")
    else:
        value = item.get("value", "")
    if item.get("required") and str(value).strip() == "":
        raise ValueError(f"Row {index}: {item.get('name', column) or 'a required field'} is empty.")
    operation = item.get("computed", "") if source == "computed" else ""
    today = date.today()
    values = {"today_day": today.day, "today_month": today.month, "today_year": today.year, "today_th_year": today.year + 543, "prev_th_year": today.year + 542, "prev_month": ((today.month + 10) % 12) + 1, "row_number": index}
    if operation == "baht_text":
        value = baht_text(value)
    elif operation in values:
        value = values[operation]
    elif operation:
        raise ValueError(f"Unsupported computed value: {operation}.")
    fmt = item.get("format", "text")
    try:
        if fmt == "number2":
            value = f"{Decimal(str(value).replace(',', '')):,.2f}"
        elif fmt == "integer":
            value = str(int(Decimal(str(value).replace(',', ''))))
        elif fmt == "uppercase":
            value = str(value).upper()
        elif fmt == "date":
            value = datetime.fromisoformat(str(value)).strftime("%d/%m/%Y")
        elif fmt == "baht_text":
            value = baht_text(value)
        elif fmt == "digits2":
            value = str(int(value)).zfill(2)
    except (ValueError, InvalidOperation):
        raise ValueError(f"Row {index}: {value!r} cannot use format {fmt}.") from None
    return str(value)


def generate(workspace: Workspace, key: str, objects: list, rows: list, *, combined=False, filename_pattern="row_{row}", preview=False):
    if not rows or len(rows) > 10000:
        raise ValueError("Choose between 1 and 10,000 data rows.")
    if not objects:
        raise ValueError("Place and bind at least one template field.")
    output = io.BytesIO()
    names = set()
    merged = fitz.open()
    try:
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            for index, row in enumerate(rows, 1):
                resolved = copy.deepcopy(objects)
                for item in resolved:
                    if item.get("type") not in {"image", "rect", "cover", "redact", "pencil"}:
                        item["value"] = resolve(item, row, index)
                pdf = edit(workspace, key, resolved)
                if preview:
                    return pdf, "preview.pdf"
                if combined:
                    with fitz.open(stream=pdf, filetype="pdf") as part:
                        merged.insert_pdf(part)
                else:
                    name = filename_pattern
                    for column, value in dict(row, row=index).items():
                        name = name.replace("{" + column + "}", str(value))
                    name = safe_name(name.removesuffix(".pdf"), f"row_{index}")
                    candidate = name + ".pdf"
                    suffix = 2
                    while candidate.casefold() in names:
                        candidate = f"{name}_{suffix}.pdf"
                        suffix += 1
                    names.add(candidate.casefold())
                    archive.writestr(candidate, pdf)
        if combined:
            return merged.tobytes(garbage=4, deflate=True), timestamp("template")
        return output.getvalue(), timestamp("template", "zip")
    finally:
        merged.close()


def migrate_legacy(workspace: Workspace, key: str, mapping: dict, coordinate_mode: str):
    if mapping.get("schema") == "pdfstudio.profile/v1":
        return mapping
    if coordinate_mode not in {"xywh", "xyxy"}:
        raise ValueError("Select the legacy coordinate convention: X/Y/width/height or X/Y/right/bottom.")
    doc = workspace.document(key)
    fields = mapping.get("fields", mapping.get("overlays", []))
    if not isinstance(fields, list) or not fields:
        raise ValueError("This mapping does not contain a fields list.")
    aliases = mapping.get("excel_mapping", {})
    globals_ = mapping.get("globals", {})
    formatters = mapping.get("formatters", {})
    unit = {"pt": 1, "mm": 72 / 25.4, "cm": 72 / 2.54}.get(mapping.get("units", "pt"))
    if unit is None:
        raise ValueError("The mapping uses unsupported units. Choose pt, mm or cm.")
    result = []
    for old in fields:
        page = int(old.get("page", 0)) + 1
        if not 1 <= page <= len(doc.pages):
            raise ValueError("A legacy field references a missing page.")
        info = doc.pages[page - 1]
        fontsize = float(old.get("fs", mapping.get("font", {}).get("size_default", 12)))
        split = old.get("split_boxes")
        offsets = None
        digits = int(old.get("digits", 0))
        if split:
            digits = int(split["digits"])
            dx = list(split.get("dx", []))
            if not dx or any(float(v) <= 0 for v in dx):
                raise ValueError("A split-digit field needs positive box spacing values.")
            while len(dx) < digits:
                dx.append(dx[-1])
            positions = [sum(float(v) for v in dx[:i]) for i in range(digits)]
            total = positions[-1] + float(dx[digits - 1])
            offsets = [v / total for v in positions]
            x, y = float(split["x0"]), float(split["y"])
            w, h = total, max(fontsize * 2 / unit, float(old.get("h", 0)))
        else:
            x, y = float(old.get("x", 0)), float(old.get("y", 0))
            w, h = float(old.get("w", 0)), float(old.get("h", 0))
            w = w if w > 0 else (x + 150 if coordinate_mode == "xyxy" else 150)
            h = h if h > 0 else (y + fontsize * 2 if coordinate_mode == "xyxy" else fontsize * 2)
            if coordinate_mode == "xyxy":
                w, h = w - x, h - y
        x, y, w, h = (v * unit for v in (x, y, w, h))
        if min(w, h) <= 0:
            raise ValueError("A legacy field has an invalid size under the selected coordinate convention.")
        source = str(split.get("value_from", "") if split else old.get("source", ""))
        datatype = str(old.get("datatype", "")).lower()
        value = old.get("datavalue", old.get("value", old.get("datafield", "")))
        kind, column, computed = "constant", "", ""
        if source.startswith("excel:") or datatype in {"excel", "excelrow", "exceldata", "excel field"}:
            alias = source.partition(":")[2] if source else str(value)
            kind, column = "excel", aliases.get(alias, alias)
        elif source.startswith("computed:") or datatype == "flex value":
            kind, computed = "computed", source.partition(":")[2] if source else str(value)
        elif source.startswith("const:"):
            value = source.partition(":")[2]
        elif datatype in {"global", "globals", "fixed value"}:
            value = globals_.get(str(value), "")
        for k, v in globals_.items():
            value = str(value).replace("{{globals." + k + "}}", str(v))
        color = split.get("color_sb", globals_.get("color_main", "#000000")) if split else old.get("color", globals_.get("color_main", "#000000"))
        from PIL import ImageColor
        try:
            color = "#" + "".join(f"{v:02x}" for v in ImageColor.getrgb(str(color))[:3])
        except ValueError:
            raise ValueError(f"Legacy field {old.get('key', '')} has an unsupported color: {color}.") from None
        # The older MuPDF overlay rotation is counterclockwise inside its box.
        rotation = int(old.get("rotation", 0)) % 360
        angle = (-rotation) % 360
        if rotation == 90:
            y += h
            w, h = h, w
        elif rotation == 180:
            x, y = x + w, y + h
        elif rotation == 270:
            x += w
            w, h = h, w
        elif rotation:
            raise ValueError("Legacy field rotations must be 0, 90, 180 or 270 degrees.")
        fmt = formatters.get(old.get("format"), old.get("format"))
        image = str(old.get("image", "")) or (str(value) if datatype == "image" else "")
        field = {"id": uuid.uuid4().hex, "type": "image" if image else "text", "name": old.get("key", "Field"), "page": page, "scope": "current", "range": str(page), "x": x / info["width"], "y": y / info["height"], "w": w / info["width"], "h": h / info["height"], "fontSize": fontsize, "font": "Sarabun", "angle": angle, "color": color, "opacity": 1, "align": "left", "sourceType": kind, "column": column, "computed": computed, "value": str(value), "format": {"num_format:#,##0.00": "number2", "num_format:00": "digits2", "baht_text": "baht_text"}.get(fmt, "text"), "digits": digits}
        if offsets:
            field["digitOffsets"] = offsets
        if image:
            field["legacyImageName"] = Path(image.replace("\\", "/")).name
            field["value"] = ""
        result.append(field)
    return {"schema": "pdfstudio.profile/v1", "coordinates": "normalized-visible-page", "name": mapping.get("template_id", "Imported mapping"), "objects": result, "migrationReviewRequired": True}
