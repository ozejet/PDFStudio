import hashlib
import io
import json
import zipfile
from pathlib import Path

import fitz
import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image

from pdfstudio import engine, templates, utilities
from pdfstudio.app import create_app
from pdfstudio.storage import Workspace


def pdf_bytes(prefix="A", pages=3):
    with fitz.open() as pdf:
        for i in range(pages):
            p = pdf.new_page(width=400, height=500)
            p.insert_text((50, 50), f"{prefix}{i+1}", fontsize=16)
            p.insert_text((80, 100), f"SECRET-{i+1}", fontsize=14)
        return pdf.tobytes()


@pytest.fixture
def workspace(tmp_path):
    value = Workspace(tmp_path / "data")
    yield value
    value.close()


def texts(data):
    with fitz.open(stream=data, filetype="pdf") as pdf:
        return [p.get_text() for p in pdf]


def obj(kind="text", **kwargs):
    return dict(id="field", type=kind, name="Test field", page=1, scope="current", x=.1, y=.3, w=.65, h=.15, angle=0, fontSize=18, font="Sarabun", align="left", color="#20365b", opacity=1, value="ADDED TEXT", sourceType="constant", column="", **kwargs)


def test_combine_reverse_repeat_rotation_and_original_preserved(workspace):
    a = workspace.import_bytes(pdf_bytes("A"), "A.pdf")
    b = workspace.import_bytes(pdf_bytes("B", 2), "B.pdf")
    original = hashlib.sha256(a.path.read_bytes()).hexdigest()
    result = engine.combine(workspace, [{"id": a.id, "label": "A", "rotations": {"3": 90}}, {"id": b.id, "label": "B"}], "A3-1,B2,A1")
    assert [t.splitlines()[0] for t in texts(result)] == ["A3", "A2", "A1", "B2", "A1"]
    with fitz.open(stream=result, filetype="pdf") as output:
        assert output[0].rotation == 90
    assert hashlib.sha256(a.path.read_bytes()).hexdigest() == original


def test_visual_selection_keeps_supplied_page_order_and_repetition(workspace):
    a = workspace.import_bytes(pdf_bytes(), "A.pdf")
    assert [t.splitlines()[0] for t in texts(engine.combine(workspace, [{"id": a.id, "pages": [3, 1, 3]}]))] == ["A3", "A1", "A3"]
    with pytest.raises(ValueError, match="Select at least"):
        engine.combine(workspace, [{"id": a.id, "pages": []}])


@pytest.mark.parametrize("value", ["0", "9", "1-x", "-1", "2--1"])
def test_invalid_ranges_rejected(value):
    with pytest.raises(ValueError):
        engine.page_range(value, 3)


def test_reverse_range_and_split(workspace):
    assert engine.page_range("3-1,2", 3, repeat=True) == [2, 1, 0, 1]
    doc = workspace.import_bytes(pdf_bytes(), "A.pdf")
    data = engine.split(workspace, doc.id, "3-2;1")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert len(archive.namelist()) == 2
        assert [t.splitlines()[0] for t in texts(archive.read(archive.namelist()[0]))] == ["A3", "A2"]


def test_image_import_and_pdf_merge(workspace):
    out = io.BytesIO()
    Image.new("RGB", (100, 70), "red").save(out, format="PNG")
    image = workspace.import_bytes(out.getvalue(), "ภาพ.png")
    doc = workspace.import_bytes(pdf_bytes(pages=1), "A.pdf")
    result = engine.combine(workspace, [{"id": image.id, "pages": [1]}, {"id": doc.id, "pages": [1]}])
    with fitz.open(stream=result, filetype="pdf") as pdf:
        assert len(pdf) == 2
        assert pdf[0].get_images()


def test_cover_retains_text_redact_removes_text_by_page_scope(workspace):
    doc = workspace.import_bytes(pdf_bytes(), "A.pdf")
    box = obj("cover")
    box.update(x=.15, y=.14, w=.5, h=.09, scope="range", range="2-3", color="#ffffff")
    assert all("SECRET" in t for t in texts(engine.edit(workspace, doc.id, [box])))
    box["type"] = "redact"
    output = texts(engine.edit(workspace, doc.id, [box]))
    assert "SECRET" in output[0]
    assert all("SECRET" not in t for t in output[1:])
    assert "SECRET" in texts(doc.path.read_bytes())[1]


def test_redaction_removes_image_pixels_and_touched_vector(workspace):
    raw = io.BytesIO()
    Image.new("RGB", (400, 500), "black").save(raw, "PNG")
    with fitz.open() as pdf:
        p = pdf.new_page(width=400, height=500)
        p.insert_image(p.rect, stream=raw.getvalue())
        p.draw_line((10, 90), (350, 90), color=(1, 0, 0), width=3)
        doc = workspace.import_bytes(pdf.tobytes(), "scanned.pdf")
    box = obj("redact")
    box.update(x=.15, y=.14, w=.5, h=.12, color="#ffffff")
    result = engine.edit(workspace, doc.id, [box])
    with fitz.open(stream=result, filetype="pdf") as pdf:
        assert not any(d["type"] == "s" for d in pdf[0].get_drawings())
        images = pdf[0].get_images()
        pixel = fitz.Pixmap(pdf, images[0][0])
        assert pixel.pixel(100, 90)[:3] == (255, 255, 255)
        assert pixel.pixel(20, 20)[:3] == (0, 0, 0)


@pytest.mark.parametrize("rotation", [0, 90, 180, 270])
def test_cropped_rotated_page_uses_visible_coordinates(workspace, rotation):
    with fitz.open() as pdf:
        p = pdf.new_page(width=400, height=500)
        p.insert_text((50, 70), "ORIGINAL")
        p.set_cropbox(fitz.Rect(15, 40, 360, 455))
        p.set_rotation(rotation)
        doc = workspace.import_bytes(pdf.tobytes(), "rotated.pdf")
    box = obj("cover")
    box.update(x=.1, y=.2, w=.3, h=.1, color="#00ff00")
    result = engine.edit(workspace, doc.id, [box])
    with fitz.open(stream=result, filetype="pdf") as pdf:
        assert pdf[0].rect.width == doc.pages[0]["width"]
        assert pdf[0].rect.height == doc.pages[0]["height"]
        pixel = pdf[0].get_pixmap()
        assert pixel.pixel(round(pixel.width*.25), round(pixel.height*.25))[:3] == (0, 255, 0)
        assert "ORIGINAL" in pdf[0].get_text()


@pytest.mark.parametrize("rotation", [90, 180, 270])
def test_rotation_normalization_preserves_visible_pixels(workspace, rotation):
    with fitz.open() as pdf:
        p = pdf.new_page(width=400, height=500)
        p.insert_text((50, 90), "ORIGINAL")
        p.draw_rect(fitz.Rect(70, 130, 190, 220), fill=(.1,.4,.7))
        p.set_cropbox(fitz.Rect(15, 40, 360, 455))
        p.set_rotation(rotation)
        before = p.get_pixmap().samples
        doc = workspace.import_bytes(pdf.tobytes(), "crop.pdf")
    with fitz.open(stream=engine.edit(workspace, doc.id, []), filetype="pdf") as output:
        assert output[0].get_pixmap().samples == before


def test_text_thai_font_image_and_pencil_exports(workspace, tmp_path):
    doc = workspace.import_bytes(pdf_bytes(pages=1), "A.pdf")
    text = obj()
    thai = obj()
    thai.update(id="thai", y=.5, value="ทดสอบ ภาษาไทย", fontSize=20, font="Sarabun Bold")
    image_bytes = io.BytesIO()
    Image.new("RGBA", (80, 30), (20, 40, 100, 150)).save(image_bytes, "PNG")
    image = workspace.add_asset(image_bytes.getvalue(), "signature.png")
    image_obj = obj("image")
    image_obj.update(id="image", x=.6, y=.75, w=.2, h=.08, assetId=image["id"])
    pencil = obj("pencil")
    pencil.update(id="pencil", points=[[.1,.8],[.2,.75],[.3,.8]], strokeWidth=3)
    result = engine.edit(workspace, doc.id, [text, thai, image_obj, pencil])
    with fitz.open(stream=result, filetype="pdf") as pdf:
        content = pdf[0].get_text()
        assert "ADDED TEXT" in content
        assert "ทดสอบ" in content
        assert pdf[0].get_images()
        assert pdf[0].get_drawings()
        pdf[0].get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(tmp_path / "edit-preview.png")


def test_redaction_removes_added_text_too(workspace):
    doc = workspace.import_bytes(pdf_bytes(pages=1), "A.pdf")
    text = obj()
    redact = obj("redact")
    redact.update(x=.05, y=.25, w=.9, h=.3)
    assert "ADDED TEXT" not in texts(engine.edit(workspace, doc.id, [text, redact]))[0]


def test_spreadsheet_preserves_padded_ids_and_sheet_selection():
    wb = Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["Name", "Tax ID", "Amount"])
    ws.append(["Example", 123, 12500])
    ws["B2"].number_format = "0000000000000"
    wb.create_sheet("Other").append(["Header"])
    data = io.BytesIO()
    wb.save(data)
    result = templates.spreadsheet(data.getvalue(), "people.xlsx", "People")
    assert result["rows"][0]["Tax ID"] == "0000000000123"
    assert result["sheets"] == ["People", "Other"]


def test_thai_baht_and_computed_values():
    assert templates.baht_text(12500) == "หนึ่งหมื่นสองพันห้าร้อยบาทถ้วน"
    assert templates.baht_text("21.25") == "ยี่สิบเอ็ดบาทยี่สิบห้าสตางค์"
    assert templates.baht_text(0) == "ศูนย์บาทถ้วน"
    item = obj()
    item.update(sourceType="computed", column="Amount", computed="baht_text")
    assert templates.resolve(item, {"Amount": 12500}) == templates.baht_text(12500)


def test_template_batch_collision_and_bindings(workspace):
    doc = workspace.import_bytes(pdf_bytes(pages=1), "A.pdf")
    item = obj()
    item.update(sourceType="excel", column="Name")
    result, name = templates.generate(workspace, doc.id, [item], [{"Name": "Alice"}, {"Name": "Bob"}], filename_pattern="employee")
    assert name.endswith(".zip")
    with zipfile.ZipFile(io.BytesIO(result)) as archive:
        assert archive.namelist() == ["employee.pdf", "employee_2.pdf"]
        assert "Alice" in texts(archive.read("employee.pdf"))[0]
        assert "Bob" in texts(archive.read("employee_2.pdf"))[0]
    with pytest.raises(ValueError, match="Missing Excel"):
        templates.generate(workspace, doc.id, [item], [{"Wrong": "Alice"}])


def test_digit_boxes_and_profile_round_trip_all_properties(workspace):
    doc = workspace.import_bytes(pdf_bytes(pages=1), "A.pdf")
    item = obj()
    item.update(digits=13, value="0123456789012", h=.08, fontSize=12, angle=90, required=True, format="text", color="#004477")
    saved = workspace.save_profile({"name": "Test", "objects": [item]}, doc)
    record, reopened = workspace.load_profile(saved["id"])
    assert record["objects"][0] == item
    assert reopened.path.read_bytes()
    item["angle"] = 0
    result = engine.edit(workspace, doc.id, [item])
    assert ''.join(texts(result)[0].split()).endswith("0123456789012")


def test_legacy_import_requires_coordinate_convention(workspace):
    doc = workspace.import_bytes(pdf_bytes(pages=1), "A.pdf")
    old = {"fields": [{"key": "name", "page": 0, "x": 40, "y": 50, "w": 200, "h": 100, "source": "excel:name"}], "excel_mapping": {"name": "Employee"}}
    with pytest.raises(ValueError, match="legacy coordinate"):
        templates.migrate_legacy(workspace, doc.id, old, "")
    result = templates.migrate_legacy(workspace, doc.id, old, "xyxy")
    assert result["objects"][0]["w"] == .4
    assert result["objects"][0]["h"] == .1
    assert result["objects"][0]["page"] == 1
    assert result["objects"][0]["column"] == "Employee"


def test_legacy_split_digits_named_colors_and_images(workspace):
    doc = workspace.import_bytes(pdf_bytes(pages=1), "A.pdf")
    old = {"fields": [{"key": "Tax ID", "page": 0, "fs": 12, "color": "Black", "split_boxes": {"digits": 4, "x0": 40, "y": 120, "dx": [20, 30, 20, 25], "value_from": "excel:tax"}}, {"key": "Signature", "page": 0, "x": 40, "y": 300, "w": 120, "h": 60, "datatype": "Image", "datavalue": "C:/old-private-folder/signature.png"}], "excel_mapping": {"tax": "Tax ID"}}
    profile = templates.migrate_legacy(workspace, doc.id, old, "xywh")
    digit, image = profile["objects"]
    assert digit["digits"] == 4
    assert digit["digitOffsets"] == [0, 20/95, 50/95, 70/95]
    assert digit["column"] == "Tax ID"
    assert image["type"] == "image"
    assert image["legacyImageName"] == "signature.png"
    assert "old-private-folder" not in json.dumps(profile)
    result, _ = templates.generate(workspace, doc.id, [digit], [{"Tax ID": "0123"}], preview=True)
    assert '0123' in ''.join(texts(result)[0].split())


def test_profile_portable_package_keeps_images(api_client):
    app, client, headers = api_client
    workspace = app.extensions['workspace']
    doc = workspace.import_bytes(pdf_bytes(pages=1), "A.pdf")
    raw = io.BytesIO()
    Image.new("RGBA", (60, 20), (20,40,80,100)).save(raw,"PNG")
    asset = workspace.add_asset(raw.getvalue(), "signature.png")
    item = obj("image")
    item["assetId"] = asset["id"]
    saved = workspace.save_profile({"name":"Portable", "objects":[item]}, doc)
    before = len(workspace.documents)
    result = client.post('/api/profiles/'+saved['id']+'/export', headers=headers, json={})
    assert len(workspace.documents) == before
    package = workspace.artifacts[result.json['id']][0].read_bytes()
    imported = client.post('/api/profiles/import',headers=headers,data={"file":(io.BytesIO(package),"profile.pdfstudio.zip")})
    assert imported.status_code == 200
    assert imported.json['profile']['objects'][0]['assetId'] == asset['id']
    assert workspace.assets[asset['id']].read_bytes() == raw.getvalue()


@pytest.mark.parametrize("name", ["word-test.docx", "slides-test.pptx"])
def test_office_import_uses_bundled_converter(workspace, name):
    sample = Path(__file__).resolve().parents[1] / "vendor/smoke-check" / name
    if not sample.exists() or not utilities.office_binary():
        pytest.skip("Prepare the bundled Office runtime and synthetic samples first")
    doc = workspace.import_bytes(sample.read_bytes(), name)
    assert doc.kind == "office"
    assert doc.pages
    assert "PDF Studio" in texts(doc.path.read_bytes())[0]


def test_table_excel_csv_and_markdown(workspace):
    with fitz.open() as pdf:
        p = pdf.new_page(width=400, height=500)
        for x in [40, 200, 360]:
            p.draw_line((x, 40), (x, 130), color=(0,0,0))
        for y in [40, 70, 100, 130]:
            p.draw_line((40,y),(360,y),color=(0,0,0))
        for x,y,t in [(50,60,"Name"),(210,60,"Amount"),(50,90,"Alice"),(210,90,"100"),(50,120,"Bob"),(210,120,"200")]:
            p.insert_text((x,y),t)
        doc = workspace.import_bytes(pdf.tobytes(), "table.pdf")
    tables = utilities.extract_tables(workspace, doc.id)
    assert tables[0]["cells"][0] == ["Name", "Amount"]
    wb = load_workbook(io.BytesIO(utilities.tables_export(tables)))
    assert wb["P1_T1"]["A2"].value == "Alice"
    assert "Alice" in utilities.markdown(workspace, doc.id)
    with zipfile.ZipFile(io.BytesIO(utilities.tables_export(tables,"csv"))) as archive:
        assert "summary.json" in archive.namelist()


def test_searchable_layer_preserves_image_and_adds_text(workspace):
    data = io.BytesIO()
    Image.new("RGB", (400, 500), "white").save(data,"PNG")
    doc = workspace.import_bytes(data.getvalue(), "scan.png")
    result = engine.searchable(workspace, doc.id, [{"page":1,"words":[{"text":"OCR RESULT","x":.1,"y":.2,"w":.3,"h":.05}]}])
    with fitz.open(stream=result,filetype="pdf") as pdf:
        assert "OCR RESULT" in pdf[0].get_text()
        pix = pdf[0].get_pixmap()
        assert pix.pixel(50,110)[:3] == (255,255,255)


def test_payment_names_handle_collisions(workspace):
    with fitz.open() as pdf:
        p = pdf.new_page()
        p.insert_text((50,50),"Transaction summary\nReference ID # 12345\nPayment\nOct 22, 2026")
        data = pdf.tobytes()
    a = workspace.import_bytes(data,"one.pdf")
    b = workspace.import_bytes(data,"two.pdf")
    names = utilities.rename_preview(workspace,[a.id,b.id],"{Date}_{ReferenceID}")
    assert [n["name"] for n in names] == ["20261022_12345.pdf","20261022_12345_2.pdf"]


@pytest.fixture
def api_client(tmp_path):
    app = create_app(tmp_path / "data",testing=True)
    client = app.test_client()
    yield app,client,{"X-Studio-Token":app.config["STUDIO_TOKEN"]}
    app.extensions["workspace"].close()


def test_api_auth_origin_and_multi_upload(api_client):
    app,client,headers = api_client
    assert client.post('/api/documents').status_code == 403
    assert client.post('/api/documents',headers={**headers,"Origin":"https://example.com"}).status_code == 403
    assert client.get('/api/health',headers={"Host":"evil.example"}).status_code == 403
    result = client.post('/api/documents',headers=headers,data={"files":[(io.BytesIO(pdf_bytes()),"A.pdf"),(io.BytesIO(b"bad pdf"),"bad.pdf")]},content_type='multipart/form-data')
    assert result.status_code == 200
    assert len(result.json["documents"]) == 1
    assert len(result.json["errors"]) == 1
    assert client.post('/api/export/combine',headers=headers,json={"items":[]}).status_code == 400


def test_native_save_cancel_and_existing_file_never_overwritten(api_client,tmp_path,monkeypatch):
    app,client,headers = api_client
    artifact = app.extensions['workspace'].artifact(pdf_bytes(),"combine_test.pdf")
    monkeypatch.setattr('pdfstudio.dialogs.choose_save',lambda *args:None)
    assert client.post('/api/save',headers=headers,json={"id":artifact['id']}).json['cancelled'] is True
    source = tmp_path/'source.pdf'
    source.write_bytes(b'ORIGINAL')
    monkeypatch.setattr('pdfstudio.dialogs.choose_save',lambda *args:str(source))
    result = client.post('/api/save',headers=headers,json={"id":artifact['id']})
    assert result.status_code == 200
    assert result.json['name'] == 'source_2.pdf'
    assert source.read_bytes() == b'ORIGINAL'
    assert (tmp_path/'source_2.pdf').read_bytes() == pdf_bytes() or (tmp_path/'source_2.pdf').stat().st_size > 100
